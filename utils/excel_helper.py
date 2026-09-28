from typing import List, Tuple
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font


def load_keywords(file_path: str, sheet_name: str, column: str,
                  start_row: int, end_row: int = 0) -> List[Tuple[int, str]]:
    """读取指定 sheet、指定列、指定行范围的内容，返回 [(行号, 关键词), ...]"""
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]

    if not end_row:
        end_row = ws.max_row

    data = []
    for row in range(start_row, end_row + 1):
        value = ws[f"{column}{row}"].value
        if value is None or str(value).strip() == "":
            continue
        data.append((row, str(value).strip()))

    wb.close()
    return data


def write_results(file_path: str, sheet_name: str, result_column: str,
                  results: List[Tuple[int, str, bool]]) -> None:
    """
    直接在原 Excel 文件上写回结果。
    PASS：正常写入文本。
    FAIL：红底 + 黑字。
    """
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    black_font = Font(color="FF000000")

    for row, _keyword, passed in results:
        cell = ws[f"{result_column}{row}"]
        if passed:
            cell.value = "PASS"
            # 清掉之前的 FAIL 样式（如果有）
            cell.fill = PatternFill(fill_type=None)
            cell.font = Font()
        else:
            cell.value = "FAIL"
            cell.fill = red_fill
            cell.font = black_font

    wb.save(file_path)